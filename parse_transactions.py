"""
Transaction Report Parser & Consolidator
========================================
Parses raw daily transaction reports from two fictional payment networks
(different file formats), enriches each record with a provider lookup
based on the terminal ID prefix, applies a fee schedule, and exports a
single consolidated, formatted Excel workbook.

Pipeline:
    1. Discover input files in ./sample_data (ALPHA_*.raw, BETA_*.txt)
    2. Parse each format with its own reader (pipe-delimited / fixed-width)
    3. Filter by response code (configurable, default "51" = insufficient funds)
    4. Look up the 4-digit terminal prefix in terminal_lookup.xlsx
       - unmatched prefixes are kept and flagged as "UNKNOWN"
    5. Apply a per-provider fee schedule
    6. Write a consolidated Excel report with a summary sheet

Usage:
    python parse_transactions.py [--rc 51] [--input sample_data]
"""

import argparse
import glob
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------

DEFAULT_INPUT_DIR = Path(__file__).parent / "sample_data"
LOOKUP_FILENAME = "terminal_lookup.xlsx"
OUTPUT_FILENAME = "consolidated_report.xlsx"

# Fee charged per declined transaction, by provider tier
FEE_SCHEDULE = {
    "PT MAJU PAYMENT": 500,
    "WARUNG DIGITAL": 750,
    "AGEN PINTAR": 750,
}
DEFAULT_FEE = 1000  # any provider not listed above
UNKNOWN_FEE = 0     # terminal prefix not found in lookup

OUTPUT_COLUMNS = [
    "SOURCE_FILE", "NETWORK", "DATE", "TIME", "TERMINAL_ID",
    "CARD_MASKED", "AMOUNT", "RESPONSE_CODE", "TRACE_NUMBER",
    "PROVIDER", "FEE",
]


@dataclass
class Transaction:
    source_file: str
    network: str
    date: datetime | None
    time: str
    terminal_id: str
    card_masked: str
    amount: int
    response_code: str
    trace_number: str
    provider: str = ""
    fee: int = 0

    def as_row(self) -> list:
        return [
            self.source_file, self.network, self.date, self.time,
            self.terminal_id, self.card_masked, self.amount,
            self.response_code, self.trace_number, self.provider, self.fee,
        ]


# ----------------------------------------------------------------------
# Lookup
# ----------------------------------------------------------------------

def load_prefix_lookup(path: Path) -> dict[str, str]:
    """
    Read terminal_lookup.xlsx into {prefix: provider}.

    The source file mimics a merged-cell layout: the provider name only
    appears on the first row of each group, so empty cells are
    forward-filled from the previous non-empty value.
    """
    if not path.exists():
        raise FileNotFoundError(f"Lookup file not found: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active

    lookup: dict[str, str] = {}
    current_provider = None
    for provider, prefix in ws.iter_rows(min_row=2, values_only=True):
        if provider is not None:
            current_provider = str(provider).strip()
        if prefix is not None and current_provider:
            lookup[str(prefix).strip()] = current_provider
    wb.close()
    return lookup


def resolve_fee(provider: str) -> int:
    if provider == "UNKNOWN":
        return UNKNOWN_FEE
    return FEE_SCHEDULE.get(provider, DEFAULT_FEE)


# ----------------------------------------------------------------------
# Format readers
# ----------------------------------------------------------------------

def parse_alpha(path: Path) -> list[Transaction]:
    """Network Alpha: pipe-delimited, 'H' header record then 'D' data records."""
    txns = []
    for line in path.read_text(encoding="utf-8").splitlines():
        parts = line.split("|")
        if parts[0] != "D" or len(parts) < 10:
            continue
        try:
            date = datetime.strptime(parts[2], "%Y%m%d")
        except ValueError:
            date = None
        t = parts[3]
        txns.append(Transaction(
            source_file=path.name,
            network="ALPHA",
            date=date,
            time=f"{t[0:2]}:{t[2:4]}:{t[4:6]}" if len(t) == 6 else t,
            terminal_id=parts[4].strip(),
            card_masked=parts[5].strip(),
            amount=int(parts[6] or 0),
            response_code=parts[7].strip(),
            trace_number=parts[8].strip(),
        ))
    return txns


# Fixed-width layout for Network Beta: (name, start, length), 1-based positions
BETA_LAYOUT = [
    ("date", 1, 6),
    ("time", 7, 6),
    ("terminal_id", 13, 8),
    ("card_masked", 21, 16),
    ("amount", 37, 13),
    ("response_code", 50, 2),
    ("trace_number", 52, 6),
]


def parse_beta(path: Path) -> list[Transaction]:
    """Network Beta: fixed-width text, first line is a report header."""
    txns = []
    lines = path.read_text(encoding="utf-8").splitlines()
    for line in lines[1:]:  # skip header line
        if len(line) < 57:
            continue
        f = {name: line[start - 1: start - 1 + length].strip()
             for name, start, length in BETA_LAYOUT}
        try:
            date = datetime.strptime(f["date"], "%y%m%d")
        except ValueError:
            date = None
        t = f["time"]
        txns.append(Transaction(
            source_file=path.name,
            network="BETA",
            date=date,
            time=f"{t[0:2]}:{t[2:4]}:{t[4:6]}" if len(t) == 6 else t,
            terminal_id=f["terminal_id"],
            card_masked=f["card_masked"],
            amount=int(f["amount"] or 0),
            response_code=f["response_code"],
            trace_number=f["trace_number"],
        ))
    return txns


READERS = {
    "ALPHA_*.raw": parse_alpha,
    "BETA_*.txt": parse_beta,
}


# ----------------------------------------------------------------------
# Excel output
# ----------------------------------------------------------------------

HEADER_FONT = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
DATA_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, columns: list[str]) -> None:
    for col, name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def write_report(txns: list[Transaction], output_path: Path) -> None:
    wb = Workbook()

    # --- Detail sheet ---
    ws = wb.active
    ws.title = "Transactions"
    style_header(ws, OUTPUT_COLUMNS)

    for r, txn in enumerate(txns, 2):
        for c, value in enumerate(txn.as_row(), 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = DATA_FONT
            cell.border = BORDER
            name = OUTPUT_COLUMNS[c - 1]
            if name == "DATE" and isinstance(value, datetime):
                cell.number_format = "DD-MMM-YY"
            elif name in ("AMOUNT", "FEE"):
                cell.number_format = "#,##0"

    widths = {"SOURCE_FILE": 22, "CARD_MASKED": 20, "TERMINAL_ID": 14,
              "PROVIDER": 22, "AMOUNT": 12, "TRACE_NUMBER": 14,
              "RESPONSE_CODE": 16, "DATE": 12}
    for c, name in enumerate(OUTPUT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 10)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{len(txns) + 1}"

    # --- Summary sheet: per provider per network ---
    summary: dict[tuple[str, str], dict] = {}
    for txn in txns:
        key = (txn.provider, txn.network)
        agg = summary.setdefault(key, {"count": 0, "amount": 0, "fee": 0})
        agg["count"] += 1
        agg["amount"] += txn.amount
        agg["fee"] += txn.fee

    ws2 = wb.create_sheet("Summary")
    cols = ["PROVIDER", "NETWORK", "TXN_COUNT", "TOTAL_AMOUNT", "TOTAL_FEE"]
    style_header(ws2, cols)
    for r, ((provider, network), agg) in enumerate(sorted(summary.items()), 2):
        for c, value in enumerate(
            [provider, network, agg["count"], agg["amount"], agg["fee"]], 1
        ):
            cell = ws2.cell(row=r, column=c, value=value)
            cell.font = DATA_FONT
            cell.border = BORDER
            if c >= 3:
                cell.number_format = "#,##0"
    for c, name in enumerate(cols, 1):
        ws2.column_dimensions[get_column_letter(c)].width = 18

    wb.save(output_path)


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Parse raw transaction reports")
    ap.add_argument("--rc", default="51",
                    help="response code to keep (default: 51). Use 'all' to keep everything")
    ap.add_argument("--input", default=str(DEFAULT_INPUT_DIR),
                    help="input directory containing raw files and the lookup")
    args = ap.parse_args()

    input_dir = Path(args.input)
    lookup = load_prefix_lookup(input_dir / LOOKUP_FILENAME)
    print(f"Loaded {len(lookup)} terminal prefixes from {LOOKUP_FILENAME}")

    all_txns: list[Transaction] = []
    for pattern, reader in READERS.items():
        files = sorted(glob.glob(str(input_dir / pattern)))
        for filepath in files:
            txns = reader(Path(filepath))
            all_txns.extend(txns)
            print(f"  {Path(filepath).name}: {len(txns):,} records")

    if not all_txns:
        print("No input files found. Run generate_sample_data.py first.")
        return

    # Filter by response code
    if args.rc.lower() != "all":
        before = len(all_txns)
        all_txns = [t for t in all_txns if t.response_code == args.rc]
        print(f"Response code filter '{args.rc}': {before:,} -> {len(all_txns):,} records")

    # Enrich: provider lookup + fee
    unmatched = 0
    for txn in all_txns:
        txn.provider = lookup.get(txn.terminal_id[:4], "UNKNOWN")
        if txn.provider == "UNKNOWN":
            unmatched += 1
        txn.fee = resolve_fee(txn.provider)

    all_txns.sort(key=lambda t: (t.date or datetime.min, t.time))

    output_path = input_dir / OUTPUT_FILENAME
    write_report(all_txns, output_path)

    total_amount = sum(t.amount for t in all_txns)
    total_fee = sum(t.fee for t in all_txns)
    print("-" * 50)
    print(f"Total records   : {len(all_txns):,}")
    print(f"Unmatched prefix: {unmatched:,} (flagged as UNKNOWN)")
    print(f"Total amount    : {total_amount:,}")
    print(f"Total fee       : {total_fee:,}")
    print(f"Report saved to : {output_path}")


if __name__ == "__main__":
    main()
